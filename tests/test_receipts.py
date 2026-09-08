import io,os,unittest,json
from unittest.mock import patch,MagicMock
from PIL import Image
import test_customer_conversation as fixture

class ReceiptTests(unittest.TestCase):
    def test_readable_png_with_damaged_metadata_is_saved_and_cleaned_for_ai(self):
        import struct,base64
        raw=self.picture().getvalue()
        # Image data is intact; a trailing text chunk has a bad checksum.
        raw=raw[:-12]+struct.pack('>I',5)+b'tEXt'+b'a\x00abc'+bytes(4)+raw[-12:]
        with Image.open(io.BytesIO(raw)) as picture:
            with self.assertRaises(SyntaxError):picture.verify()
        uploaded=self.client.post('/receipts',data={'receipt':(io.BytesIO(raw),'Screenshot.PNG')},content_type='multipart/form-data')
        self.assertEqual(uploaded.status_code,302)
        receipt=self.mod.q('SELECT * FROM purchase_receipts',one=True)
        with open(os.path.join(self.mod.app.config['UPLOAD_FOLDER'],'receipts',receipt['filename']),'rb') as original:
            self.assertEqual(original.read(),raw)
        result=dict(supplier='Shop',receipt_date='',items='Detergent',total='12',vat='',currency='GBP',notes='')
        response=MagicMock();response.__enter__.return_value.read.return_value=json.dumps({'output_text':json.dumps(result)}).encode()
        with patch.dict(os.environ,{'OPENAI_API_KEY':'test'}),patch.object(self.mod.urllib.request,'urlopen',return_value=response) as api:
            self.mod.receipt_ai_extract(receipt)
        image_url=json.loads(api.call_args.args[0].data)['input'][0]['content'][1]['image_url']
        self.assertTrue(image_url.startswith('data:image/jpeg;base64,'))
        with Image.open(io.BytesIO(base64.b64decode(image_url.split(',')[1]))) as picture:
            picture.load();self.assertEqual(picture.size,(20,20))

    def test_truncated_jpeg_pixels_are_rejected(self):
        data=io.BytesIO();Image.new('RGB',(200,200),'white').save(data,'JPEG')
        response=self.client.post('/receipts',data={'receipt':(io.BytesIO(data.getvalue()[:-100]),'photo.jpg')},content_type='multipart/form-data')
        self.assertEqual(response.status_code,200)
        self.assertEqual(self.mod.q('SELECT count(*) n FROM purchase_receipts',one=True)['n'],0)

    def setUp(self):
        fixture.CustomerConversationTests.setUp(self)
        self.mod.app.config['UPLOAD_FOLDER']=os.path.join(self.tmp.name,'uploads')
    tearDown=fixture.CustomerConversationTests.tearDown
    def picture(self):
        data=io.BytesIO();Image.new('RGB',(20,20),'white').save(data,'PNG');data.seek(0);return data
    def upload(self):
        with patch.object(self.mod,'receipt_ai_extract',side_effect=RuntimeError('Enter details manually.')):
            return self.client.post('/receipts',data={'receipt':(self.picture(),'receipt.png')},content_type='multipart/form-data')
    def test_upload_duplicate_review_save_and_private_original(self):
        first=self.upload();self.assertEqual(first.status_code,302)
        second=self.upload();self.assertEqual(first.location.split('?')[0],second.location.split('?')[0])
        self.assertEqual(self.mod.q('SELECT count(*) n FROM purchase_receipts',one=True)['n'],1)
        rid=self.mod.q('SELECT id FROM purchase_receipts',one=True)['id']
        self.assertEqual(self.mod.q('SELECT count(*) n FROM expenses',one=True)['n'],0)
        self.assertEqual(self.mod.app.test_client().get(f'/receipts/{rid}/file').status_code,302)
        original=self.client.get(f'/receipts/{rid}/file');self.assertEqual(original.mimetype,'image/png');original.close()
        values={'supplier':'Cleaning supplier','receipt_date':'2026-09-07','items':'Detergent','total':'24.00','vat':'4.00','currency':'GBP','notes':'Checked'}
        self.client.post(first.location,data=values);self.client.post(first.location,data=values)
        self.assertEqual(self.mod.q('SELECT count(*) n FROM expenses',one=True)['n'],1)
        expense=self.mod.q('SELECT * FROM expenses',one=True)
        self.assertEqual(expense['amount'],20)
        self.assertEqual(expense['vat_amount'],4)
    def test_uncertain_amount_is_not_saved_and_invalid_file_rejected(self):
        url=self.upload().location
        response=self.client.post(url,data={'supplier':'Shop','receipt_date':'2026-09-07','total':'NaN','vat':'','currency':'GBP'})
        self.assertEqual(response.status_code,200)
        self.assertEqual(self.mod.q('SELECT count(*) n FROM expenses',one=True)['n'],0)
        self.client.post('/receipts',data={'receipt':(io.BytesIO(b'<html>bad</html>'),'receipt.jpg')},content_type='multipart/form-data')
        self.assertEqual(self.mod.q('SELECT count(*) n FROM purchase_receipts',one=True)['n'],1)
    def test_ai_extraction_keeps_unknowns_for_review(self):
        self.upload();receipt=self.mod.q('SELECT * FROM purchase_receipts',one=True)
        result=dict(supplier='Shop',receipt_date='',items='Cleaning liquid',total='12.00',vat='',currency='GBP',notes='Date and VAT unreadable')
        response=MagicMock();response.__enter__.return_value.read.return_value=json.dumps({'output_text':json.dumps(result)}).encode()
        with patch.dict(os.environ,{'OPENAI_API_KEY':'test'}),patch.object(self.mod.urllib.request,'urlopen',return_value=response) as api:
            self.mod.receipt_ai_extract(receipt)
        payload=json.loads(api.call_args.args[0].data)
        self.assertFalse(payload['store']);self.assertEqual(payload['input'][0]['content'][1]['type'],'input_image')
        saved=self.mod.q('SELECT * FROM purchase_receipts',one=True)
        self.assertEqual(saved['vat'],'');self.assertEqual(saved['status'],'Needs review')
        self.assertIsNone(saved['expense_id'])

    def test_upload_returns_before_ai_and_export_is_safe(self):
        with patch.object(self.mod,'receipt_ai_extract') as ai:
            result=self.client.post('/receipts',data={'receipt':(self.picture(),'receipt.png')},content_type='multipart/form-data')
            self.assertEqual(result.status_code,302);ai.assert_not_called()
        self.mod.run("UPDATE purchase_receipts SET supplier='=FORMULA',total='20' ")
        result=self.client.get('/receipts/export.csv')
        self.assertEqual(result.status_code,200)
        self.assertIn(b"'=FORMULA",result.data)
        self.assertIn(b'Original file',result.data)

    def test_excel_export_has_clickable_receipt_and_summary(self):
        self.upload();receipt=self.mod.q('SELECT * FROM purchase_receipts',one=True)
        self.mod.run("UPDATE purchase_receipts SET supplier='Supplier',category='Materials & chemicals',total='24.00',vat='4.00' WHERE id=?",(receipt['id'],))
        result=self.client.get('/receipts/export.xlsx')
        self.assertEqual(result.status_code,200)
        self.assertIn('spreadsheetml',result.mimetype)
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(result.data))
        self.assertEqual(book.sheetnames,['Receipts','Summary'])
        self.assertEqual(book['Receipts']['I2'].value,'Open receipt')
        self.assertIn(f'/receipts/{receipt["id"]}/file',book['Receipts']['I2'].hyperlink.target)

    def test_manual_cash_expense_without_receipt(self):
        result=self.client.post('/expenses',data={'expense_date':'2026-09-07','category':'Fuel & travel','supplier':'Petrol station','description':'Diesel','amount':'50','vat_amount':'','payment_method':'Cash'})
        self.assertEqual(result.status_code,302)
        expense=self.mod.q('SELECT * FROM expenses',one=True)
        self.assertEqual(expense['payment_method'],'Cash')
        self.assertEqual(expense['category'],'Fuel & travel')
        page=self.client.get('/expenses')
        self.assertEqual(page.status_code,200)
        self.assertIn(b'Estimated profit',page.data)
        self.assertIn(b'Paid by cash',page.data)

    def test_iphone_receipt_upload_is_accepted(self):
        from pillow_heif import register_heif_opener
        register_heif_opener()
        data=io.BytesIO();Image.new('RGB',(64,64),'white').save(data,format='HEIF');data.seek(0)
        result=self.client.post('/receipts',data={'receipt':(data,'IMG_001.HEIC')},content_type='multipart/form-data')
        self.assertEqual(result.status_code,302)
        self.assertEqual(self.mod.q('SELECT mime_type FROM purchase_receipts',one=True)['mime_type'],'image/heic')

    def test_iphone_screenshot_with_misleading_filename_is_accepted(self):
        result=self.client.post('/receipts',data={'receipt':(self.picture(),'Screenshot 2026-09-07.HEIC')},content_type='multipart/form-data')
        self.assertEqual(result.status_code,302)
        self.assertEqual(self.mod.q('SELECT mime_type FROM purchase_receipts',one=True)['mime_type'],'image/png')

    def test_browser_converted_screenshot_is_normalised_to_png(self):
        data=io.BytesIO();Image.new('RGB',(64,64),'white').save(data,format='TIFF');data.seek(0)
        result=self.client.post('/receipts',data={'receipt':(data,'Screenshot.tiff')},content_type='multipart/form-data')
        self.assertEqual(result.status_code,302)
        receipt=self.mod.q('SELECT * FROM purchase_receipts',one=True)
        self.assertEqual(receipt['mime_type'],'image/png')
        self.assertTrue(receipt['filename'].endswith('.png'))

    def test_background_failure_keeps_original_and_shows_retry(self):
        self.upload();receipt=self.mod.q('SELECT * FROM purchase_receipts',one=True)
        self.mod.run("UPDATE purchase_receipts SET status='Reading' WHERE id=?",(receipt['id'],))
        with patch.object(self.mod,'receipt_ai_extract',side_effect=RuntimeError('AI connection unavailable')):
            self.mod.read_receipt_background(receipt['id'])
        updated=self.mod.q('SELECT * FROM purchase_receipts',one=True)
        self.assertEqual(updated['status'],'Needs review')
        self.assertEqual(updated['ai_error'],'AI connection unavailable')
        self.assertTrue(os.path.isfile(os.path.join(self.mod.app.config['UPLOAD_FOLDER'],'receipts',updated['filename'])))

    def test_duplicate_unread_receipt_retries_but_preserves_reviewed_fields(self):
        self.upload()
        self.assertIn('read=1',self.upload().location)
        self.mod.run("UPDATE purchase_receipts SET supplier='My correction'")
        self.assertNotIn('read=1',self.upload().location)

    def test_empty_ai_result_is_visible_failure(self):
        self.upload();receipt=self.mod.q('SELECT * FROM purchase_receipts',one=True)
        response=MagicMock();response.__enter__.return_value.read.return_value=json.dumps({'output_text':json.dumps(dict.fromkeys(('supplier','receipt_date','items','total','vat','currency','notes'),''))}).encode()
        with patch.dict(os.environ,{'OPENAI_API_KEY':'test'}),patch.object(self.mod.urllib.request,'urlopen',return_value=response):
            self.mod.read_receipt_background(receipt['id'])
        page=self.client.get('/receipts/'+str(receipt['id']))
        self.assertIn(b'No receipt details could be read',page.data)
        self.assertIn(b'Receipt reading failed',page.data)
        self.assertEqual(self.mod.q('SELECT count(*) n FROM expenses',one=True)['n'],0)

    def test_reading_endpoint_persists_and_displays_extracted_details(self):
        self.upload();receipt=self.mod.q('SELECT * FROM purchase_receipts',one=True)
        result=dict(supplier='Receipt test supplier',receipt_date='2026-09-07',items='Carpet detergent',total='24.00',vat='4.00',currency='GBP',notes='')
        response=MagicMock();response.__enter__.return_value.read.return_value=json.dumps({'output_text':json.dumps(result)}).encode()
        endpoint='/receipts/'+str(receipt['id'])+'/reading'
        with patch.dict(os.environ,{'OPENAI_API_KEY':'test'}),patch.object(self.mod.urllib.request,'urlopen',return_value=response):
            started=self.client.post(endpoint)
            self.assertEqual(started.status_code,200)
            self.mod.receipt_read_threads[receipt['id']].join(timeout=5)
        self.assertFalse(self.client.get(endpoint).json['reading'])
        page=self.client.get('/receipts/'+str(receipt['id']))
        for value in ('Receipt test supplier','Carpet detergent','24.00','4.00','Saved receipt details'):
            self.assertIn(value.encode(),page.data)
        self.assertEqual(self.mod.q('SELECT count(*) n FROM expenses',one=True)['n'],0)
